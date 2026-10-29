"""
FAZ 4 — Gold etiketlerini yeni kurallarla (sürüm 2) gözden geçir

İki adım:

  make   data/gold/gozden_gecirme.xlsx dosyasını oluşturur: 40 gold ürünün her biri
         için tek satır (sayfa metni, numaralı negatif yorumlar, mevcut etiketler ve
         not) ve yanında BOŞ "yeni_..." sütunları. Öneri / tahmin yazmaz.

  apply  Doldurulan "yeni_..." sütunlarını data/gold/etiketleme.xlsx (Etiketler)
         dosyasına işler, değişenleri (eski -> yeni) listeler ve
         data/gold/gozden_gecirme_degisiklikler.csv'ye yazar, ardından
         src/03_export_gold.py'yi çalıştırır.
         --dry-run: sadece değişiklikleri gösterir, hiçbir dosyaya yazmaz.

Doldurma kuralı (gozden_gecirme.xlsx > Nasıl doldurulur sayfası):
  - Bir satırdaki tüm "yeni_" hücreleri boşsa o ürünün etiketi DEĞİŞMEZ.
  - Bir satırda herhangi bir "yeni_" hücresi doluysa o satır YENİ ETİKETİN TAMAMIDIR:
    yeni_gap_var ve yeni_seviye zorunlu; konu sütunlarında X olmayanlar işaretsiz sayılır.
    yeni_not boşsa mevcut not korunur.

Kullanım:
    python src/04_review_gold.py make
    python src/04_review_gold.py apply --dry-run
    python src/04_review_gold.py apply
"""
import argparse
import csv
import importlib.util
import json
import subprocess
import sys
from pathlib import Path

import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
GOLD_DIR = ROOT / "data" / "gold"
LABELS = GOLD_DIR / "etiketleme.xlsx"
REVIEW = GOLD_DIR / "gozden_gecirme.xlsx"
CHANGES = GOLD_DIR / "gozden_gecirme_degisiklikler.csv"
SHEET = "Gözden Geçirme"

GAP_TYPES = ["Beden/Kalıp", "Renk", "Görsel", "Kalite", "Malzeme", "Dayanıklılık", "Eksik Bilgi"]
LEVELS = ["Yok", "Orta", "Yüksek"]
# Etiketler sayfasındaki sütun sırası: no, parent_asin, Sayfada yazan, Negatif yorumlar,
# veri_uygun, gap_var, <7 konu>, seviye, not
COL_GAP, COL_T0, COL_LVL, COL_NOTE = 6, 7, 14, 15

INFO = ["no", "parent_asin", "baslik", "sayfa (özellikler + açıklama)", "negatif_yorumlar"]
CURRENT = ["mevcut_gap_var"] + [f"mevcut_{t}" for t in GAP_TYPES] + ["mevcut_seviye", "mevcut_not"]
NEW = ["yeni_gap_var"] + [f"yeni_{t}" for t in GAP_TYPES] + ["yeni_seviye", "yeni_not"]

HOWTO = [
    "GÖZDEN GEÇİRME — NASIL DOLDURULUR",
    "",
    "Her satırda bir gold ürün var. Kurallar: docs/etiketleme_kurallari.md (sürüm 2) ve etiketleme.xlsx > Talimatlar.",
    "Soldaki sütunlar bilgi içindir (sayfa metni, yorumlar, MEVCUT etiketin). Sağdaki sarı 'yeni_' sütunları BOŞ gelir.",
    "",
    "1) Etiket DEĞİŞMİYORSA: o satırın 'yeni_' hücrelerinin HEPSİNİ boş bırak.",
    "2) Etiket DEĞİŞİYORSA: o satırda yeni etiketin TAMAMINI yaz:",
    "   - yeni_gap_var ve yeni_seviye zorunlu",
    "   - yeni konu sütunları: işaretli olması gereken HER konuya X koy (X olmayan konu işaretsiz sayılır)",
    "   - yeni_not: yeni gerekçe. Boş bırakırsan mevcut not korunur.",
    "3) Bu dosyadaki diğer sütunları değiştirme.",
    "",
    "Bitince: python src/04_review_gold.py apply --dry-run   (önce değişiklikleri gör)",
    "         python src/04_review_gold.py apply             (etiketleme.xlsx'e işle + gold_40.csv'yi yeniden üret)",
]


def load_sheet_module():
    spec = importlib.util.spec_from_file_location("labeling", ROOT / "src" / "02_make_labeling_sheet.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def clean(v):
    return None if v is None or str(v).strip() == "" else str(v).strip()


def read_labels():
    """etiketleme.xlsx > Etiketler: parent_asin -> (satır no, mevcut etiket)."""
    wb = load_workbook(LABELS)
    ws = wb["Etiketler"]
    out = {}
    for row in range(2, ws.max_row + 1):
        asin = clean(ws.cell(row, 2).value)
        if not asin:
            continue
        out[asin] = {
            "row": row,
            "no": ws.cell(row, 1).value,
            "gap_var": clean(ws.cell(row, COL_GAP).value),
            "types": {t: "X" if clean(ws.cell(row, COL_T0 + i).value) else None for i, t in enumerate(GAP_TYPES)},
            "seviye": clean(ws.cell(row, COL_LVL).value),
            "not": clean(ws.cell(row, COL_NOTE).value),
        }
    return wb, out


# ---------------------------------------------------------------- make
def make():
    if REVIEW.exists():
        raise SystemExit(f"{REVIEW} zaten var; doldurduklarının üzerine yazmamak için durdum. "
                         "Yeniden oluşturmak istiyorsan önce dosyayı sil.")
    cfg = yaml.safe_load(open(ROOT / "config.yaml", encoding="utf-8"))
    products = {p["parent_asin"]: p for p in map(json.loads, open(
        ROOT / "data" / "processed" / f"products_{cfg['n_products']}.jsonl", encoding="utf-8"))}
    _, labels = read_labels()

    wb = Workbook()
    how = wb.active
    how.title = "Nasıl doldurulur"
    for i, line in enumerate(HOWTO, 1):
        how.cell(row=i, column=1, value=line)
    how["A1"].font = Font(bold=True, size=14)
    how.column_dimensions["A"].width = 120
    rules = wb.create_sheet("Kurallar (sürüm 2)")
    load_sheet_module().write_instructions(rules)

    ws = wb.create_sheet(SHEET)
    headers = INFO + CURRENT + NEW
    ws.append(headers)
    for asin, lab in sorted(labels.items(), key=lambda kv: kv[1]["no"]):
        p = products[asin]
        page = "\n".join(f"• {f}" for f in p["features"])
        if p["description"]:
            page += "\n\nAÇIKLAMA: " + p["description"]
        reviews = "\n\n".join(f"{i}) [{int(r['rating'])}★] {r['title']} — {r['text']}"
                              for i, r in enumerate(p["negative_reviews"], 1))
        ws.append([lab["no"], asin, p["title"], page, reviews,
                   lab["gap_var"], *[lab["types"][t] for t in GAP_TYPES], lab["seviye"], lab["not"]]
                  + [None] * len(NEW))

    n = ws.max_row
    head_fill = PatternFill("solid", fgColor="1F3A5F")
    cur_fill = PatternFill("solid", fgColor="E7E6E6")
    new_fill = PatternFill("solid", fgColor="FFF4CC")
    c_cur, c_new = len(INFO) + 1, len(INFO) + len(CURRENT) + 1
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    for r in range(2, n + 1):
        ws.row_dimensions[r].height = 300
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=r, column=c)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if c_cur <= c < c_new:
                cell.fill = cur_fill
            elif c >= c_new:
                cell.fill = new_fill
    widths = {1: 5, 2: 13, 3: 30, 4: 60, 5: 80}
    for c in range(1, len(headers) + 1):
        w = widths.get(c, 10)
        if headers[c - 1] in ("mevcut_not", "yeni_not"):
            w = 45
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "D2"

    def dv(options, col_from, col_to):
        v = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
        ws.add_data_validation(v)
        v.add(f"{get_column_letter(col_from)}2:{get_column_letter(col_to)}{n}")

    dv(["Evet", "Hayır"], c_new, c_new)
    dv(["X"], c_new + 1, c_new + len(GAP_TYPES))
    dv(LEVELS, c_new + len(GAP_TYPES) + 1, c_new + len(GAP_TYPES) + 1)
    wb.active = wb.sheetnames.index(SHEET)
    wb.save(REVIEW)
    print(f"Yazıldı: {REVIEW}  ({n - 1} ürün)")


# ---------------------------------------------------------------- apply
def read_review():
    ws = load_workbook(REVIEW, read_only=True)[SHEET]
    rows = list(ws.iter_rows(values_only=True))
    col = {h: i for i, h in enumerate(rows[0])}
    missing = [h for h in INFO + CURRENT + NEW if h not in col]
    if missing:
        raise SystemExit(f"gozden_gecirme.xlsx sütunları eksik/değişmiş: {missing}")
    return [{h: r[col[h]] for h in INFO + CURRENT + NEW} for r in rows[1:] if any(v is not None for v in r)]


def apply(dry_run):
    wb, labels = read_labels()
    review = read_review()
    errors, changes, updates = [], [], []

    asins = [clean(r["parent_asin"]) for r in review]
    if sorted(asins) != sorted(labels):
        errors.append("gozden_gecirme.xlsx'teki ASIN'ler etiketleme.xlsx > Etiketler ile aynı değil")

    for r in review:
        asin = clean(r["parent_asin"])
        lab = labels.get(asin)
        if lab is None:
            continue
        where = f"#{r['no']} ({asin})"
        # gözden geçirme dosyası oluşturulduktan sonra etiketleme.xlsx değişmiş mi?
        cur = {"gap_var": clean(r["mevcut_gap_var"]), "seviye": clean(r["mevcut_seviye"]),
               "not": clean(r["mevcut_not"]),
               "types": {t: "X" if clean(r[f"mevcut_{t}"]) else None for t in GAP_TYPES}}
        if cur != {k: lab[k] for k in cur}:
            errors.append(f"{where}: 'mevcut_' sütunları etiketleme.xlsx ile uyuşmuyor "
                          "(dosya sonradan değişmiş olabilir; gözden geçirme dosyasını yeniden üret)")
            continue

        if all(clean(r[h]) is None for h in NEW):
            continue  # değişiklik yok
        gap, lvl = clean(r["yeni_gap_var"]), clean(r["yeni_seviye"])
        note = clean(r["yeni_not"]) or lab["not"]
        types = {}
        for t in GAP_TYPES:
            v = clean(r[f"yeni_{t}"])
            if v is not None and v.upper() != "X":
                errors.append(f"{where}: yeni_{t} hücresinde X dışı değer ({v!r})")
            types[t] = "X" if v else None
        n_types = sum(1 for v in types.values() if v)
        if gap not in ("Evet", "Hayır"):
            errors.append(f"{where}: yeni_gap_var boş/geçersiz ({gap!r})")
        if lvl not in LEVELS:
            errors.append(f"{where}: yeni_seviye boş/geçersiz ({lvl!r})")
        if gap == "Hayır" and (n_types or lvl != "Yok"):
            errors.append(f"{where}: yeni_gap_var = Hayır ise konu işaretsiz ve seviye Yok olmalı")
        if gap == "Evet" and (not n_types or lvl not in ("Orta", "Yüksek")):
            errors.append(f"{where}: yeni_gap_var = Evet ise en az 1 konu ve seviye Orta/Yüksek olmalı")

        new = {"gap_var": gap, "types": types, "seviye": lvl, "not": note}
        diff = []
        if gap != lab["gap_var"]:
            diff.append(("gap_var", lab["gap_var"], gap))
        for t in GAP_TYPES:
            if types[t] != lab["types"][t]:
                diff.append((t, "X" if lab["types"][t] else "-", "X" if types[t] else "-"))
        if lvl != lab["seviye"]:
            diff.append(("seviye", lab["seviye"], lvl))
        if note != lab["not"]:
            diff.append(("not", lab["not"], note))
        if diff:
            updates.append((lab, new))
            changes += [(lab["no"], asin, f, o, nv) for f, o, nv in diff]

    if errors:
        print(f"HATA ({len(errors)}) — hiçbir şey yazılmadı:")
        for e in errors:
            print(f"  - {e}")
        raise SystemExit(1)

    print(f"Değişen ürün: {len(updates)} / {len(labels)}  |  değişen alan: {len(changes)}")
    for no, asin, field, old, new in changes:
        if field == "not":
            print(f"  #{no:<2} {asin}  not:\n        eski: {old}\n        yeni: {new}")
        else:
            print(f"  #{no:<2} {asin}  {field}: {old} -> {new}")
    if dry_run:
        print("\n--dry-run: hiçbir dosyaya yazılmadı.")
        return
    if not updates:
        print("Değişiklik yok; etiketleme.xlsx'e dokunulmadı.")
    else:
        ws = wb["Etiketler"]
        for lab, new in updates:
            # .value ile yazılır: ws.cell(r, c, None) hücreyi temizlemez (X kaldırılamazdı)
            ws.cell(lab["row"], COL_GAP).value = new["gap_var"]
            for i, t in enumerate(GAP_TYPES):
                ws.cell(lab["row"], COL_T0 + i).value = new["types"][t]
            ws.cell(lab["row"], COL_LVL).value = new["seviye"]
            ws.cell(lab["row"], COL_NOTE).value = new["not"]
        wb.save(LABELS)
        print(f"\nYazıldı: {LABELS}")
    with open(CHANGES, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(["no", "parent_asin", "alan", "eski", "yeni"])
        w.writerows(changes)
    print(f"Değişiklik listesi: {CHANGES}")

    print("\n--- src/03_export_gold.py ---", flush=True)
    res = subprocess.run([sys.executable, str(ROOT / "src" / "03_export_gold.py")])
    if res.returncode:
        raise SystemExit(res.returncode)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    sub.add_parser("make", help="gozden_gecirme.xlsx oluştur")
    p = sub.add_parser("apply", help="yeni etiketleri etiketleme.xlsx'e işle")
    p.add_argument("--dry-run", action="store_true", help="sadece değişiklikleri göster")
    args = ap.parse_args()
    make() if args.cmd == "make" else apply(args.dry_run)


if __name__ == "__main__":
    main()
