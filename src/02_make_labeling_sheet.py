"""
FAZ 4 — Etiketleme dosyasını hazırla

products_200.jsonl içindeki 40 "gold" ürünü bir Excel dosyasına yazar.
Sen sadece sağdaki sarı sütunları doldurursun.

Çıktı: data/gold/etiketleme.xlsx

Kullanım:
    python src/02_make_labeling_sheet.py
"""
import json
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]

GAP_TYPES = ["Beden/Kalıp", "Renk", "Görsel", "Kalite", "Malzeme", "Dayanıklılık", "Eksik Bilgi"]


def page_text(p):
    parts = [f"BAŞLIK: {p['title']}"]
    if p.get("price") is not None:
        parts.append(f"FİYAT: ${p['price']}")
    if p.get("features"):
        parts.append("ÖZELLİKLER:\n" + "\n".join(f"• {x}" for x in p["features"]))
    if p.get("description"):
        parts.append("AÇIKLAMA: " + p["description"][:800])
    if p.get("details"):
        d = "; ".join(f"{k}: {v}" for k, v in list(p["details"].items())[:8])
        parts.append("DETAYLAR: " + d)
    return "\n\n".join(parts)


def reviews_text(p):
    lines = []
    for i, r in enumerate(p["negative_reviews"], 1):
        t = r["text"][:500]
        lines.append(f"{i}) [{int(r['rating'])}★] {r['title']} — {t}")
    return "\n\n".join(lines)


def main():
    cfg = yaml.safe_load(open(ROOT / "config.yaml", encoding="utf-8"))
    src = ROOT / "data" / "processed" / f"products_{cfg['n_products']}.jsonl"
    products = [json.loads(l) for l in open(src, encoding="utf-8")]
    gold = [p for p in products if p["split"] == "gold"]

    wb = Workbook()

    # ---------- Talimat sayfası ----------
    ins = wb.active
    ins.title = "Talimatlar"
    rules = [
        "ETİKETLEME TALİMATI",
        "",
        "Her satırda bir ürün var. Ürünü GÖRMENE gerek yok; sadece 'Sayfada yazan' ve 'Negatif yorumlar' sütunlarını oku.",
        "Soru: Sayfanın verdiği söz ile müşterilerin şikayeti arasında çelişki var mı?",
        "",
        "0) veri_uygun: Önce bunu doldur. Aşağıdakilerden biri varsa 'Hayır' yaz, notu yaz ve o satırın geri kalanını BOŞ bırak:",
        "   - Yorumlar sayfadaki ürünle ilgili değil (başka ürünler anlatılıyor; Amazon farklı ürünleri birleştirmiş)",
        "   - Ürün giyim/ayakkabı değil (filtreye rağmen kaçan)",
        "   - Geriye değerlendirilebilecek 2'den az yorum kalıyor",
        "   'Hayır' çıkan her gold ürün için 'Yedekler' sayfasındaki SIRADAKİ ürünü etiketle (atlama yapma, sırayla git).",
        "   Toplamda veri_uygun = Evet olan 40 ürün olmalı.",
        "",
        "1) gap_var: Evet / Hayır",
        "2) Konu sütunları: çelişki hangi konudaysa o sütuna X koy (birden fazla olabilir).",
        "   Kural: Bir konu ancak EN AZ 2 yorum o sorundan bahsediyorsa işaretlenir.",
        "3) seviye:",
        "   Yüksek = Sayfa açıkça bir şey iddia etmiş, yorumlar tersini söylüyor (ör. 'su geçirmez' yazıyor, 'su aldı' diyorlar)",
        "   Orta   = Sayfa o konuda bilgi vermemiş, yorumlar şikayet ediyor (ör. beden bilgisi yok, 'küçük geldi' diyorlar)",
        "   Yok    = Çelişki yok (gap_var = Hayır ise seviye de Yok)",
        "4) not: Neden bu kararı verdin, 1 cümle.",
        "",
        "SAYILMAYANLAR: kargo/paketleme, satıcı/iade süreci, 'beğenmedim' gibi kişisel zevk, tek seferlik arıza (1 yorum).",
        "TEKRARLAR: Birebir aynı yorumlar script tarafından zaten ayıklandı. Yine de aynı kişinin neredeyse aynı yorumunu görürsen tek say.",
        "",
        "ÖRNEK: Sayfa 'Waterproof leather boot' diyor; 2 yorum 'su aldı', 1 yorum 'kargo geç'.",
        "        -> gap_var = Evet | Malzeme = X (su geçirmezlik malzeme iddiası) | seviye = Yüksek | not: 'Su geçirmez diyor, 2/2 ilgili yorum su aldı diyor'",
        "",
        "ÖNEMLİ: Bu dosyayı bitirip GitHub'a commit etmeden yapay zeka sistemini ÇALIŞTIRMA.",
    ]
    for i, r in enumerate(rules, 1):
        ins.cell(row=i, column=1, value=r)
    ins["A1"].font = Font(bold=True, size=14)
    ins.column_dimensions["A"].width = 130

    reserve = [p for p in products if p["split"] == "reserve"]
    build_sheet(wb, "Etiketler", gold)
    if reserve:
        build_sheet(wb, "Yedekler", reserve)

    out = ROOT / "data" / "gold" / "etiketleme.xlsx"
    wb.save(out)
    print(f"Yazıldı: {out}  ({len(gold)} gold + {len(reserve)} yedek)")


def build_sheet(wb, title, items):
    from openpyxl.utils import get_column_letter
    ws = wb.create_sheet(title)
    headers = (["no", "parent_asin", "Sayfada yazan", "Negatif yorumlar", "veri_uygun", "gap_var"]
               + GAP_TYPES + ["seviye", "not"])
    ws.append(headers)
    for i, p in enumerate(items, 1):
        ws.append([i, p["parent_asin"], page_text(p), reviews_text(p), None, None]
                  + [None] * len(GAP_TYPES) + [None, None])

    head_fill = PatternFill("solid", fgColor="1F3A5F")
    input_fill = PatternFill("solid", fgColor="FFF4CC")
    check_fill = PatternFill("solid", fgColor="E2EFDA")
    n = len(items) + 1
    col_ok, col_gap = 5, 6
    col_t0, col_t1 = 7, 6 + len(GAP_TYPES)
    col_lvl, col_note = col_t1 + 1, col_t1 + 2
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    for r in range(2, n + 1):
        ws.row_dimensions[r].height = 300
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=r, column=c)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if c == col_ok:
                cell.fill = check_fill
            elif c >= col_gap:
                cell.fill = input_fill

    for col, w in {"A": 5, "B": 13, "C": 55, "D": 75}.items():
        ws.column_dimensions[col].width = w
    ws.column_dimensions[get_column_letter(col_ok)].width = 10
    ws.column_dimensions[get_column_letter(col_gap)].width = 9
    for c in range(col_t0, col_t1 + 1):
        ws.column_dimensions[get_column_letter(c)].width = 9
    ws.column_dimensions[get_column_letter(col_lvl)].width = 10
    ws.column_dimensions[get_column_letter(col_note)].width = 35
    ws.freeze_panes = "C2"

    def dv(options, col_from, col_to):
        v = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
        ws.add_data_validation(v)
        v.add(f"{get_column_letter(col_from)}2:{get_column_letter(col_to)}{n}")

    dv(["Evet", "Hayır"], col_ok, col_ok)
    dv(["Evet", "Hayır"], col_gap, col_gap)
    dv(["X"], col_t0, col_t1)
    dv(["Yok", "Orta", "Yüksek"], col_lvl, col_lvl)


if __name__ == "__main__":
    main()
