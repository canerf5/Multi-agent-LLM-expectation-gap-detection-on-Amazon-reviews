"""
FAZ 9 — Alt Soru 3: önerilerin sayfa bağlamıyla tutarlılığının kör puanlanması (kapsam §6; yontem §7.5)

İki adım:

  make   Kalan (rest) kümeden seed = 42 ile 20 ürün seçer; her ürün için sayfa metni, numaralı
         yorumlar ve iki sistemin önerilerini "Sistem A" / "Sistem B" olarak yazar. Hangi sistemin
         A olduğu ürün başına rastgele (aynı tohumla) belirlenir ve YALNIZCA anahtar dosyasına yazılır.
         Çıktılar (değerlendirme klasöründe):
           oneri_puanlama.xlsx     puanlama dosyası (rubrik, sarı puan sütunları, açılır liste)
           oneri_anahtar.json      A/B eşlemesi — puanlama bitene kadar AÇILMAZ
         Dosyalar varsa üzerine yazmaz.

  ozet   Doldurulmuş puanları anahtarla eşler; oneri_puanlama.csv (ürün, sistem, puan) ve
         oneri_puanlama_ozet.md yazar. Puanlar eksikse durur.

Kullanım:
    python src/05_alt_soru3.py make
    python src/05_alt_soru3.py ozet
"""
import argparse
import csv
import json
import random
import sys
from pathlib import Path
from statistics import mean, median

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

import agents
import pipeline

ROOT = Path(__file__).resolve().parents[1]
EVAL_DIR = ROOT / "reports" / "degerlendirme" / "2026-10-30_215304"
RUNS = {"multi_agent": ROOT / "outputs" / "runs" / "all" / "2026-10-30_204205_multi_agent_faz7-tam-kosu",
        "baseline": ROOT / "outputs" / "runs" / "all" / "2026-10-30_212510_baseline_faz7-tam-kosu"}
RESULT_FILE = {"multi_agent": "agent3.json", "baseline": "baseline.json"}
N, SEED = 20, 42
SHEET, NO_SUGGESTION = "Puanlama", "yok"
SCORES = ["1", "2", "3", "4", "5"]

RUBRIC = [
    ("5", "Öneriler sayfadaki ilgili ifadeye veya eksikliğe ve yorumlardaki şikayete doğrudan dayanıyor; "
          "düzeltme somut ve uygulanabilir; uydurma ayrıntı yok"),
    ("4", "Öneriler sayfa ve şikayetle tutarlı ve uygulanabilir; küçük bir eksik var (ör. düzeltme içeriği biraz genel)"),
    ("3", "Öneriler ilgili ama genel veya belirsiz (\"bilgi ekle\" düzeyinde) ya da bir kısmı dayanaksız"),
    ("2", "Sayfa veya şikayetle bağ zayıf; sayfada olmayan bir iddiaya dayanıyor ya da şikayetlerde geçmeyen ayrıntı içeriyor"),
    ("1", "Öneriler ilgisiz, yanlış ya da sayfa yerine ürünün kendisini değiştirmeyi öneriyor"),
]
HOWTO = [
    "ALT SORU 3 — ÖNERİLERİN SAYFA BAĞLAMIYLA TUTARLILIĞI (KÖR PUANLAMA)",
    "",
    "Her satırda bir ürün var: sayfa metni, numaralı negatif yorumlar ve iki sistemin önerileri (Sistem A, Sistem B).",
    "Hangi sistemin A, hangisinin B olduğu ürün başına rastgele belirlendi ve gizli tutuluyor. Tahmin etmeye çalışma.",
    "",
    "Soru: Bu sistemin önerileri, bu ürünün sayfası ve yorumları bağlamında tutarlı mı?",
    "Her sistemin önerilerinin TAMAMINA tek puan ver (1–5). Puanı sarı hücreye açılır listeden seç; istersen not yaz.",
    f"Bir sistem o üründe gap bulmadıysa (öneri yoksa) hücre gri gelir ve \"{NO_SUGGESTION}\" yazar; puan verme.",
    "",
    "RUBRİK",
    *[f"  {p} = {t}" for p, t in RUBRIC],
    "",
    "Bitince: python src/05_alt_soru3.py ozet",
]


def select_products():
    """Kalan kümeden (dosya sırasıyla) seed = 42 ile 20 ürün; aynı üreteçle ürün başına A/B ataması."""
    rest = pipeline.load_products("rest")
    rng = random.Random(SEED)
    chosen = rng.sample(rest, N)
    key = {p["parent_asin"]: dict(zip("AB", rng.sample(list(RUNS), 2))) for p in chosen}
    return chosen, key


def suggestions_text(res):
    if not res["oneriler"]:
        return None
    lines = []
    for i, o in enumerate(res["oneriler"], 1):
        lines += [f"{i}) [{o['tip']} | {o['islem']}] {o['duzeltme']}",
                  f"   sayfa ifadesi: {o['sayfa_iddiasi']!r}",
                  f"   müşteri: {o['musteri_deneyimi']}"]
    return "\n".join(lines)


def make(out_dir=EVAL_DIR):
    xlsx, key_file = out_dir / "oneri_puanlama.xlsx", out_dir / "oneri_anahtar.json"
    for f in (xlsx, key_file):
        if f.exists():
            raise SystemExit(f"{f} zaten var; puanların üzerine yazmamak için durdum.")
    chosen, key = select_products()

    wb = Workbook()
    how = wb.active
    how.title = "Nasıl puanlanır"
    for i, line in enumerate(HOWTO, 1):
        how.cell(row=i, column=1, value=line)
    how["A1"].font = Font(bold=True, size=14)
    how.column_dimensions["A"].width = 140

    ws = wb.create_sheet(SHEET)
    headers = ["no", "parent_asin", "Sayfa", "Negatif yorumlar", "Sistem A önerileri", "Puan A",
               "Sistem B önerileri", "Puan B", "Not"]
    ws.append(headers)
    grey, yellow = PatternFill("solid", fgColor="D9D9D9"), PatternFill("solid", fgColor="FFF4CC")
    dv = DataValidation(type="list", formula1='"' + ",".join(SCORES) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    for i, p in enumerate(chosen, 1):
        texts = {}
        for side in "AB":
            res = json.loads((RUNS[key[p["parent_asin"]][side]] / p["parent_asin"] /
                              RESULT_FILE[key[p["parent_asin"]][side]]).read_text(encoding="utf-8"))
            texts[side] = suggestions_text(res)
        ws.append([i, p["parent_asin"], agents.page_text(p), agents.reviews_text(p),
                   texts["A"] or "(Bu sistem bu üründe gap bulmadı; öneri yok)", None if texts["A"] else NO_SUGGESTION,
                   texts["B"] or "(Bu sistem bu üründe gap bulmadı; öneri yok)", None if texts["B"] else NO_SUGGESTION,
                   None])
        r = i + 1
        for col, side in ((6, "A"), (8, "B")):
            cell = ws.cell(row=r, column=col)
            if texts[side]:
                cell.fill = yellow
                dv.add(cell.coordinate)
            else:
                cell.fill = grey
        ws.cell(row=r, column=9).fill = yellow
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3A5F")
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    for r in range(2, len(chosen) + 2):
        ws.row_dimensions[r].height = 320
        for c in range(1, len(headers) + 1):
            ws.cell(row=r, column=c).alignment = Alignment(wrap_text=True, vertical="top")
    for col, w in {"A": 5, "B": 13, "C": 55, "D": 70, "E": 55, "F": 8, "G": 55, "H": 8, "I": 30}.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "C2"
    wb.active = wb.sheetnames.index(SHEET)
    out_dir.mkdir(parents=True, exist_ok=True)
    wb.save(xlsx)
    key_file.write_text(json.dumps({"tohum": SEED, "urun": N, "kume": "rest",
                                    "kosular": {s: d.name for s, d in RUNS.items()}, "eslesme": key},
                                   ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Yazıldı: {xlsx}\nAnahtar (puanlama bitene kadar açılmaz): {key_file}")


def ozet(out_dir=EVAL_DIR):
    ws = load_workbook(out_dir / "oneri_puanlama.xlsx", read_only=True)[SHEET]
    rows = list(ws.iter_rows(values_only=True))
    col = {h: i for i, h in enumerate(rows[0])}
    key = json.loads((out_dir / "oneri_anahtar.json").read_text(encoding="utf-8"))["eslesme"]
    out, missing = [], []
    for r in rows[1:]:
        if r[col["parent_asin"]] is None:
            continue
        asin = r[col["parent_asin"]]
        for side in "AB":
            v = r[col[f"Puan {side}"]]
            v = None if v is None else str(v).strip()
            if v == NO_SUGGESTION:
                out.append({"parent_asin": asin, "sistem": key[asin][side], "taraf": side, "puan": "", "not": "öneri yok"})
            elif v in SCORES:
                out.append({"parent_asin": asin, "sistem": key[asin][side], "taraf": side, "puan": int(v),
                            "not": r[col["Not"]] or ""})
            else:
                missing.append(f"{asin} Puan {side} ({v!r})")
    if missing:
        raise SystemExit("Eksik/geçersiz puan — özet yazılmadı:\n  " + "\n  ".join(missing))
    with open(out_dir / "oneri_puanlama.csv", "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(out[0]))
        w.writeheader()
        w.writerows(out)
    L = ["# Alt Soru 3 — Öneri puanlaması", "", f"{len(key)} ürün (kalan küme, tohum {SEED}), kör puanlama, rubrik 1–5.", "",
         "| Sistem | Puanlanan ürün | Öneri yok | Ortalama | Medyan | 1 | 2 | 3 | 4 | 5 |", "|---|---|---|---|---|---|---|---|---|---|"]
    for s in RUNS:
        sc = [o["puan"] for o in out if o["sistem"] == s and o["puan"] != ""]
        none = sum(o["sistem"] == s and o["puan"] == "" for o in out)
        dist = [sum(x == k for x in sc) for k in range(1, 6)]
        L.append(f"| {s} | {len(sc)} | {none} | {mean(sc):.2f} | {median(sc):.1f} | " + " | ".join(map(str, dist)) + " |"
                 if sc else f"| {s} | 0 | {none} | — | — | | | | | |")
    (out_dir / "oneri_puanlama_ozet.md").write_text("\n".join(L) + "\n", encoding="utf-8")
    print("\n".join(L))


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("cmd", choices=["make", "ozet"])
    args = ap.parse_args(argv)
    make() if args.cmd == "make" else ozet()


if __name__ == "__main__":
    sys.exit(main())
