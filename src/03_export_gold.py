"""
FAZ 4 — Gold etiketleri dışa aktar

data/gold/etiketleme.xlsx dosyasındaki elle etiketleri okur, doğrular ve
değerlendirme script'lerinin kullanacağı makine okunur bir CSV'ye yazar.

Ürün seçimi (Talimatlar sayfasındaki kural):
  - Etiketler sayfasında veri_uygun = Evet olan ürünler alınır.
  - veri_uygun = Hayır olan her gold ürün için Yedekler sayfasındaki SIRADAKİ
    ürün alınır (atlama yok; yedek de Hayır ise bir sonrakine geçilir).
  - Kullanılmayan yedekler CSV'ye girmez.

Doğrulama (hata varsa yazmadan durur):
  - tam n_gold (40) satır, parent_asin tekrarı yok
  - boş hücre yok (veri_uygun, gap_var, seviye, not)
  - gap_var = 0 ise tüm tipler 0 ve seviye Yok
  - gap_var = 1 ise en az bir tip 1 ve seviye Orta / Yüksek

Çıktı: data/gold/gold_40.csv

Kullanım:
    python src/03_export_gold.py
"""
import csv
import json
from pathlib import Path

import yaml
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]

GAP_TYPES = ["Beden/Kalıp", "Renk", "Görsel", "Kalite", "Malzeme", "Dayanıklılık", "Eksik Bilgi"]
TYPE_COLS = ["beden_kalip", "renk", "gorsel", "kalite", "malzeme", "dayaniklilik", "eksik_bilgi"]
LEVELS = ["Yok", "Orta", "Yüksek"]


def clean(v):
    return None if v is None or str(v).strip() == "" else str(v).strip()


def read_sheet(ws, source):
    """Sayfadaki satırları sözlük olarak döndürür; geçersiz değerleri hata listesine yazar."""
    head = [c.value for c in ws[1]]
    col = {h: i for i, h in enumerate(head)}
    rows, errors = [], []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if all(v is None for v in r):
            continue
        no, asin = r[col["no"]], clean(r[col["parent_asin"]])
        where = f"{source} #{no} ({asin})"
        rec = {"no": no, "parent_asin": asin, "kaynak": source, "where": where,
               "veri_uygun": clean(r[col["veri_uygun"]]), "not": clean(r[col["not"]])}
        if rec["veri_uygun"] not in ("Evet", "Hayır"):
            errors.append(f"{where}: veri_uygun boş/geçersiz ({rec['veri_uygun']!r})")
        rec["gap_raw"] = clean(r[col["gap_var"]])
        rec["seviye"] = clean(r[col["seviye"]])
        rec["types"] = {}
        for t, c in zip(GAP_TYPES, TYPE_COLS):
            v = clean(r[col[t]])
            if v is not None and v.upper() != "X":
                errors.append(f"{where}: {t} hücresinde X dışı değer ({v!r})")
            rec["types"][c] = 1 if v else 0
        rows.append(rec)
    return rows, errors


def validate(rec):
    """Kullanılacak (veri_uygun = Evet) bir satırın tutarlılık hataları."""
    w, errs = rec["where"], []
    for name in ("parent_asin", "not"):
        if not rec[name]:
            errs.append(f"{w}: {name} boş")
    if rec["gap_raw"] not in ("Evet", "Hayır"):
        errs.append(f"{w}: gap_var boş/geçersiz ({rec['gap_raw']!r})")
        return errs
    if rec["seviye"] not in LEVELS:
        errs.append(f"{w}: seviye boş/geçersiz ({rec['seviye']!r})")
    n_types = sum(rec["types"].values())
    if rec["gap_raw"] == "Hayır":
        if n_types:
            errs.append(f"{w}: gap_var = Hayır ama {n_types} tip işaretli")
        if rec["seviye"] != "Yok":
            errs.append(f"{w}: gap_var = Hayır ama seviye {rec['seviye']!r} (Yok olmalı)")
    else:
        if not n_types:
            errs.append(f"{w}: gap_var = Evet ama hiç tip işaretli değil")
        if rec["seviye"] not in ("Orta", "Yüksek"):
            errs.append(f"{w}: gap_var = Evet ama seviye {rec['seviye']!r} (Orta/Yüksek olmalı)")
    return errs


def main():
    cfg = yaml.safe_load(open(ROOT / "config.yaml", encoding="utf-8"))
    n_gold = cfg["n_gold"]
    wb = load_workbook(ROOT / "data" / "gold" / "etiketleme.xlsx", read_only=True)
    gold, e1 = read_sheet(wb["Etiketler"], "gold")
    reserve, e2 = read_sheet(wb["Yedekler"], "yedek")
    errors = e1 + e2

    # Örneklemle eşleşme: ASIN'ler products_200.jsonl'deki gold/reserve kümesinde olmalı
    products = [json.loads(l) for l in open(ROOT / "data" / "processed" / f"products_{cfg['n_products']}.jsonl",
                                            encoding="utf-8")]
    split_of = {p["parent_asin"]: p["split"] for p in products}
    for rec in gold + reserve:
        expected = "gold" if rec["kaynak"] == "gold" else "reserve"
        if split_of.get(rec["parent_asin"]) != expected:
            errors.append(f"{rec['where']}: örneklemde split = {expected} değil "
                          f"({split_of.get(rec['parent_asin'])!r})")

    # Seçim: gold Evet + her gold Hayır için sıradaki uygun yedek
    excluded = [r for r in gold if r["veri_uygun"] == "Hayır"]
    selected = [r for r in gold if r["veri_uygun"] == "Evet"]
    need, used = len(excluded), []
    for r in reserve:
        if len(used) == need:
            break
        if r["veri_uygun"] == "Hayır":
            excluded.append(r)
            continue
        if r["veri_uygun"] == "Evet":
            used.append(r)
    if len(used) < need:
        errors.append(f"Yedek yetmedi: {need} gold ürün çıkarıldı, {len(used)} uygun yedek var")
    selected += used

    for rec in selected:
        errors += validate(rec)
    asins = [r["parent_asin"] for r in selected]
    if len(set(asins)) != len(asins):
        errors.append("parent_asin tekrarı var")
    if len(selected) != n_gold:
        errors.append(f"Satır sayısı {len(selected)}, beklenen {n_gold}")

    print(f"Gold veri_uygun: Evet {sum(r['veri_uygun'] == 'Evet' for r in gold)}, "
          f"Hayır {sum(r['veri_uygun'] == 'Hayır' for r in gold)} | kullanılan yedek: {len(used)}")
    print(f"\nveri_uygun = Hayır olan ürünler ({len(excluded)}):")
    for r in excluded:
        print(f"  - {r['where']}: {r['not'] or '(not yok)'}")
    if not excluded:
        print("  (yok)")

    if errors:
        print(f"\nHATA ({len(errors)}) — CSV yazılmadı:")
        for e in errors:
            print(f"  - {e}")
        raise SystemExit(1)

    out = ROOT / "data" / "gold" / f"gold_{n_gold}.csv"
    with open(out, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(["parent_asin", "gap_var"] + TYPE_COLS + ["seviye", "not", "kaynak"])
        for r in selected:
            w.writerow([r["parent_asin"], 1 if r["gap_raw"] == "Evet" else 0]
                       + [r["types"][c] for c in TYPE_COLS] + [r["seviye"], r["not"], r["kaynak"]])
    print(f"\nDoğrulama geçti. Yazıldı: {out}  ({len(selected)} satır: "
          f"{len(selected) - len(used)} gold + {len(used)} yedek)")


if __name__ == "__main__":
    main()
